import tempfile
import unittest
from pathlib import Path

import database as db
from openpyxl import load_workbook
from werkzeug.security import check_password_hash


class DatabaseTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.old_path = db.DATABASE_PATH
        db.DATABASE_PATH = Path(self.temp.name) / "test.xlsx"
        db.initialize_database()

    def tearDown(self):
        db.DATABASE_PATH = self.old_path
        self.temp.cleanup()

    def test_workbook_schema_and_seed_hash_once(self):
        wb = load_workbook(db.DATABASE_PATH)
        self.assertEqual(wb.sheetnames, ["Users", "Books", "Feedback"])
        self.assertEqual([c.value for c in wb["Users"][1]], db.USERS_HEADERS)
        admin = db.get_user_by_email(" ADMIN@gmail.com ")
        self.assertTrue(check_password_hash(admin["Password_Hash"], "Admin@123"))
        wb.close()
        db.initialize_database()
        self.assertEqual(len([u for u in db.get_all_users() if u["Role"] == "Admin"]), 1)

    def test_users_books_feedback(self):
        student = db.add_user("Sam Student", "  SAM@example.com ", "Pass123")
        self.assertEqual(db.get_user_by_email("sam@example.com")["User_ID"], student)
        with self.assertRaises(db.DatabaseError): db.add_user("Duplicate", "sam@example.com", "Pass123")
        book = db.add_book("Python", "A. Author", "Computing", "Intro", "AUTO")
        self.assertEqual(db.search_books("author")[0]["Book_ID"], book)
        feedback_id = db.add_feedback(student, "Please add more books")
        db.reply_to_feedback(feedback_id, "We will add them soon")
        self.assertEqual(db.get_feedback_by_user(student)[0]["Status"], "Replied")
        db.delete_user(student)
        self.assertEqual(len(db.get_all_feedback()), 1)


if __name__ == "__main__": unittest.main()
