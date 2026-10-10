"""Small Excel-based data helpers for the library project."""
import os
import threading
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from werkzeug.security import generate_password_hash

BASE_DIR = Path(__file__).resolve().parent
DATABASE_PATH = Path(os.environ.get("LIBRARY_DATABASE", BASE_DIR / "library_database.xlsx"))
USERS_HEADERS = ["User_ID", "Name", "Email", "Password_Hash", "Role", "Profile_Image"]
BOOKS_HEADERS = ["Book_ID", "Title", "Author", "Category", "Description", "PDF_Path", "Date_Added"]
FEEDBACK_HEADERS = ["Feedback_ID", "User_ID", "Message", "Admin_Reply", "Date_Submitted", "Status"]
_write_lock = threading.RLock()


class DatabaseError(Exception):
    """A readable database operation error."""


def _open():
    if not DATABASE_PATH.exists():
        initialize_database()
    try:
        return load_workbook(DATABASE_PATH)
    except Exception as exc:
        raise DatabaseError("The Excel workbook could not be opened. Close it in Excel and try again.") from exc


def _rows(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(v is not None for v in row)]


def _next_id(sheet, prefix):
    numbers = []
    for row in sheet.iter_rows(min_row=2, min_col=1, max_col=1, values_only=True):
        value = row[0]
        if isinstance(value, str) and value.startswith(prefix) and value[len(prefix):].isdigit():
            numbers.append(int(value[len(prefix):]))
    return f"{prefix}{max(numbers, default=0) + 1:03d}"


def initialize_database():
    """Create missing sheets/headers without clearing existing records; seed admin once."""
    with _write_lock:
        DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)
        if DATABASE_PATH.exists():
            try:
                book = load_workbook(DATABASE_PATH)
            except Exception as exc:
                raise DatabaseError("The existing Excel workbook is damaged or unavailable.") from exc
        else:
            book = Workbook()
            book.active.title = "Users"
            book.create_sheet("Books")
            book.create_sheet("Feedback")
        expected = {"Users": USERS_HEADERS, "Books": BOOKS_HEADERS, "Feedback": FEEDBACK_HEADERS}
        for name, headers in expected.items():
            if name not in book.sheetnames:
                book.create_sheet(name)
            sheet = book[name]
            second_row = [sheet.cell(2, column).value for column in range(1, len(headers) + 1)]
            first_row = [sheet.cell(1, column).value for column in range(1, len(headers) + 1)]
            # Repair workbooks made by an early version that appended headers below the blank first row.
            if second_row == headers:
                if all(value is None for value in first_row):
                    sheet.delete_rows(1)
                elif first_row == headers:
                    sheet.delete_rows(2)
            if all(sheet.cell(1, column).value is None for column in range(1, len(headers) + 1)):
                for column, header in enumerate(headers, start=1):
                    sheet.cell(1, column).value = header
        # Retain unrelated sheets in a supplied workbook; newly created workbooks have exactly three.
        users = book["Users"]
        headers = [cell.value for cell in users[1]]
        if all(header in headers for header in USERS_HEADERS):
            email_col = headers.index("Email") + 1
            admin_email = os.environ.get("INITIAL_ADMIN_EMAIL", "admin@gmail.com").strip().lower()
            exists = any(str(users.cell(r, email_col).value or "").strip().lower() == admin_email for r in range(2, users.max_row + 1))
            if not exists:
                password = os.environ.get("INITIAL_ADMIN_PASSWORD", "Admin@123")
                values = {"User_ID": _next_id(users, "U"), "Name": "Library Administrator", "Email": admin_email,
                          "Password_Hash": generate_password_hash(password), "Role": "Admin", "Profile_Image": "static/profiles/default.svg"}
                users.append([values[h] for h in headers])
        book.save(DATABASE_PATH)
        book.close()


def _read(sheet_name):
    book = _open()
    try:
        if sheet_name not in book.sheetnames:
            raise DatabaseError(f"The workbook is missing the {sheet_name} worksheet.")
        return _rows(book[sheet_name])
    finally:
        book.close()


def _write(action):
    with _write_lock:
        book = _open()
        try:
            result = action(book)
            book.save(DATABASE_PATH)
            return result
        except DatabaseError:
            raise
        except Exception as exc:
            raise DatabaseError("The workbook could not be updated. Close it in Excel and try again.") from exc
        finally:
            book.close()


def get_all_users(): return _read("Users")
def get_user_by_email(email):
    normalized = (email or "").strip().lower()
    return next((u for u in get_all_users() if str(u.get("Email") or "").strip().lower() == normalized), None)
def get_user_by_id(user_id): return next((u for u in get_all_users() if u.get("User_ID") == user_id), None)
def get_all_books(): return _read("Books")
def get_book_by_id(book_id): return next((b for b in get_all_books() if b.get("Book_ID") == book_id), None)
def search_books(search_text):
    query = (search_text or "").strip().casefold()
    books = get_all_books()
    return [b for b in books if not query or any(query in str(b.get(k) or "").casefold() for k in ("Title", "Author", "Category"))]
def get_all_feedback(): return _read("Feedback")
def get_feedback_by_user(user_id): return [f for f in get_all_feedback() if f.get("User_ID") == user_id]


def add_user(name, email, password, role="User"):
    name, email = (name or "").strip(), (email or "").strip().lower()
    if not name or not email or not password or role not in ("Admin", "User"):
        raise DatabaseError("Name, email and password are required.")
    def action(book):
        sheet = book["Users"]
        headers = [c.value for c in sheet[1]]
        email_col = headers.index("Email") + 1
        if any(str(sheet.cell(row, email_col).value or "").strip().lower() == email for row in range(2, sheet.max_row + 1)):
            raise DatabaseError("An account with that email already exists.")
        values = {"User_ID": _next_id(sheet, "U"), "Name": name, "Email": email,
                  "Password_Hash": generate_password_hash(password), "Role": role, "Profile_Image": "static/profiles/default.svg"}
        sheet.append([values.get(h) for h in headers])
        return values["User_ID"]
    return _write(action)


def delete_user(user_id):
    def action(book):
        sheet = book["Users"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, 1).value == user_id:
                if sheet.cell(row, 5).value == "Admin": raise DatabaseError("Administrator accounts cannot be deleted here.")
                sheet.delete_rows(row)
                return True
        return False
    return _write(action)


def update_user_password(user_id, password):
    if not password: raise DatabaseError("A new password is required.")
    def action(book):
        sheet = book["Users"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, 1).value == user_id:
                sheet.cell(row, 4).value = generate_password_hash(password)
                return True
        return False
    return _write(action)


def add_book(title, author, category, description, pdf_path):
    values = [title, author, category, pdf_path]
    if any(not str(v or "").strip() for v in values): raise DatabaseError("Title, author, category and PDF are required.")
    def action(book):
        sheet = book["Books"]
        book_id = _next_id(sheet, "B")
        stored_path = f"books/{book_id}.pdf" if pdf_path == "AUTO" else pdf_path
        sheet.append([book_id, title.strip(), author.strip(), category.strip(), (description or "").strip(), stored_path, datetime.now().strftime("%Y-%m-%d")])
        return book_id
    return _write(action)


def update_book_pdf_path(book_id, pdf_path):
    def action(book):
        sheet = book["Books"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, 1).value == book_id:
                sheet.cell(row, 6).value = pdf_path
                return True
        return False
    return _write(action)


def delete_book(book_id):
    def action(book):
        sheet = book["Books"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, 1).value == book_id:
                sheet.delete_rows(row)
                return True
        return False
    return _write(action)


def add_feedback(user_id, message):
    message = (message or "").strip()
    if not message: raise DatabaseError("Feedback message cannot be empty.")
    def action(book):
        sheet = book["Feedback"]
        feedback_id = _next_id(sheet, "F")
        sheet.append([feedback_id, user_id, message, "", datetime.now().strftime("%Y-%m-%d %H:%M"), "Pending"])
        return feedback_id
    return _write(action)


def reply_to_feedback(feedback_id, reply):
    reply = (reply or "").strip()
    if not reply: raise DatabaseError("Reply cannot be empty.")
    def action(book):
        sheet = book["Feedback"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, 1).value == feedback_id:
                sheet.cell(row, 4).value = reply
                sheet.cell(row, 6).value = "Replied"
                return True
        return False
    return _write(action)
